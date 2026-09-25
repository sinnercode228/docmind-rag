"""Test helpers: tiny in-memory PDF/DOCX builders and SSE parsing."""

from __future__ import annotations

import io
import json
from typing import Any


def make_pdf(pages: list[str], title: str | None = None) -> bytes:
    """Build a minimal, valid PDF with one text line per page (Helvetica)."""
    objects: list[bytes] = []
    n_pages = len(pages)
    # 1: catalog, 2: pages, 3: font, then (page, content) pairs, then optional info
    kids = " ".join(f"{4 + 2 * i} 0 R" for i in range(n_pages))
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {n_pages} >>".encode())
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for i, text in enumerate(pages):
        escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode()
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {5 + 2 * i} 0 R >>".encode()
        )
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    info_ref = b""
    if title:
        objects.append(f"<< /Title ({title}) >>".encode())
        info_ref = f" /Info {len(objects)} 0 R".encode()

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R".encode()
        + info_ref
        + f" >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return out.getvalue()


def make_docx(title: str, sections: dict[str, list[str]]) -> bytes:
    import docx

    document = docx.Document()
    document.add_heading(title, level=0)
    for heading, paragraphs in sections.items():
        document.add_heading(heading, level=1)
        for paragraph in paragraphs:
            document.add_paragraph(paragraph)
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Budget"
    table.rows[0].cells[1].text = "600 EUR"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def parse_sse(body: str) -> list[tuple[str, dict[str, Any]]]:
    events = []
    for block in body.strip().split("\n\n"):
        event, data = "message", []
        for line in block.splitlines():
            if line.startswith("event:"):
                event = line[6:].strip()
            elif line.startswith("data:"):
                data.append(line[5:].strip())
        if data:
            events.append((event, json.loads("\n".join(data))))
    return events
